import os
from pptx import Presentation

pptx_path = r"D:\7TH SEM MAJOR F\presentation.pptx"
assert os.path.exists(pptx_path), "File does not exist!"
prs = Presentation(pptx_path)

print(f"Presentation validated! Total slides: {len(prs.slides)}")
print(f"Dimensions: {prs.slide_width.inches:.2f} x {prs.slide_height.inches:.2f} inches (16:9 widescreen)")

keywords_to_verify = [
    "18,200", "12,740", "2,730", "97.40%", "Model V1", "Model V3",
    "Grad-CAM", "COVID-19", "Lung Opacity", "Normal", "Pneumonia",
    "Tuberculosis", "675 ms", "825 ms", "1.50 seconds", "Nine exact",
    "academic prototype"
]

all_text = ""
for i, slide in enumerate(prs.slides, 1):
    slide_texts = []
    has_image = False
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                slide_texts.append(p.text)
        if shape.shape_type == 13: # Picture
            has_image = True
    
    combined = " | ".join([t for t in slide_texts if t.strip()])
    all_text += " " + combined
    img_tag = " [Has Picture]" if has_image else ""
    print(f"Slide {i:02d}: {combined[:75]}...{img_tag}")

print("\n--- Keyword & Accuracy Verification ---")
for kw in keywords_to_verify:
    found = kw.lower() in all_text.lower()
    print(f"  [{'PASS' if found else 'FAIL'}] {kw}")

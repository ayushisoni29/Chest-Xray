import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # blank layout

    # Color Palette
    NAVY_DARK = RGBColor(15, 35, 65)     # #0F2341
    SLATE_BLUE = RGBColor(30, 75, 130)   # #1E4B82
    TEAL_ACCENT = RGBColor(14, 165, 160) # #0EA5A0
    BG_LIGHT = RGBColor(248, 250, 252)   # #F8FAFC
    CARD_BG = RGBColor(255, 255, 255)    # #FFFFFF
    TEXT_DARK = RGBColor(30, 41, 59)     # #1E293B
    TEXT_MUTED = RGBColor(100, 116, 139) # #64748B
    BORDER_COLOR = RGBColor(226, 232, 240)
    ALERT_BG = RGBColor(254, 243, 199)   # #FEF3C7
    ALERT_BORDER = RGBColor(245, 158, 11)# #F59E0B
    ALERT_TEXT = RGBColor(146, 64, 14)

    def set_slide_background(slide, color=BG_LIGHT):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background() # no border
        return bg

    def add_header(slide, title_text, category_text):
        # Header banner bar
        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.9))
        banner.fill.background()
        banner.line.fill.background()
        tf = banner.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = TEAL_ACCENT
        p_cat.font.name = "Arial"
        
        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY_DARK
        p_title.font.name = "Arial"

        # Footer
        footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.3))
        footer.fill.background()
        footer.line.fill.background()
        ft_tf = footer.text_frame
        ft_tf.margin_left = ft_tf.margin_top = ft_tf.margin_right = ft_tf.margin_bottom = 0
        p_ft = ft_tf.paragraphs[0]
        p_ft.text = "Chest X-Ray Multi-Disease Detection & Grad-CAM Explainability | B.Tech Major Project"
        p_ft.font.size = Pt(9)
        p_ft.font.color.rgb = TEXT_MUTED
        p_ft.font.name = "Arial"

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_COLOR):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    # ==================== SLIDE 1: Title Slide ====================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, NAVY_DARK)
    
    # Title decorative accent line
    accent = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.4), Inches(1.5), Inches(0.08))
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL_ACCENT
    accent.line.fill.background()

    tb_title = s1.shapes.add_textbox(Inches(1.2), Inches(1.6), Inches(11.0), Inches(2.2))
    tf1 = tb_title.text_frame
    tf1.word_wrap = True
    
    p1 = tf1.paragraphs[0]
    p1.text = "Chest X-Ray Multi-Disease Detection &\nGrad-CAM Explainability System"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)
    p1.font.name = "Arial"

    p2 = tf1.add_paragraph()
    p2.text = "A Dual-Model ResNet50 Deep Learning Ensemble with Visual Saliency Explanations for Thoracic Radiography"
    p2.font.size = Pt(15)
    p2.font.color.rgb = RGBColor(203, 213, 225) # Slate 300
    p2.space_before = Pt(10)
    p2.font.name = "Arial"

    # Meta box
    meta_card = add_card(s1, Inches(1.2), Inches(4.3), Inches(10.9), Inches(2.2), bg_color=RGBColor(24, 48, 82), border_color=RGBColor(45, 75, 115))
    tf_meta = meta_card.text_frame
    tf_meta.word_wrap = True
    tf_meta.margin_left = Inches(0.3)
    tf_meta.margin_top = Inches(0.25)
    
    pm1 = tf_meta.paragraphs[0]
    pm1.text = "B.Tech Final Year Major Project Examination / Viva Voce"
    pm1.font.size = Pt(13)
    pm1.font.bold = True
    pm1.font.color.rgb = TEAL_ACCENT
    pm1.font.name = "Arial"

    pm2 = tf_meta.add_paragraph()
    pm2.text = "• Student Name / Roll No:  [Student Name — Roll Number / ID] (Editable Placeholder)\n" \
               "• Project Supervisor / Guide:  [Supervisor / Faculty Advisor Name] (Editable Placeholder)\n" \
               "• Department:  Department of Computer Science & Engineering\n" \
               "• Institution:  [College / University Name] (Editable Placeholder)\n" \
               "• Academic Session:  2025 – 2026"
    pm2.font.size = Pt(11)
    pm2.font.color.rgb = RGBColor(241, 245, 249)
    pm2.space_before = Pt(6)
    pm2.font.name = "Arial"

    # ==================== SLIDE 2: Problem Statement & Motivation ====================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Clinical Problem Statement & Motivation", "Context & Healthcare Drivers")

    c1 = add_card(s2, Inches(0.8), Inches(1.6), Inches(3.7), Inches(5.0))
    tf_c1 = c1.text_frame
    tf_c1.word_wrap = True
    tf_c1.margin_left = tf_c1.margin_right = Inches(0.25)
    tf_c1.margin_top = Inches(0.3)
    p = tf_c1.paragraphs[0]
    p.text = "High Disease Burden"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = NAVY_DARK
    p = tf_c1.add_paragraph()
    p.text = "• Thoracic conditions (COVID-19, Pneumonia, Tuberculosis, diffuse lung opacities) are leading global causes of mortality.\n\n" \
             "• Chest X-ray (CXR) is the universal frontline diagnostic modality due to low cost, rapid acquisition, and high availability.\n\n" \
             "• Early triage is essential to prevent critical pulmonary complications."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(12)

    c2 = add_card(s2, Inches(4.8), Inches(1.6), Inches(3.7), Inches(5.0))
    tf_c2 = c2.text_frame
    tf_c2.word_wrap = True
    tf_c2.margin_left = tf_c2.margin_right = Inches(0.25)
    tf_c2.margin_top = Inches(0.3)
    p = tf_c2.paragraphs[0]
    p.text = "Diagnostic Bottlenecks"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = NAVY_DARK
    p = tf_c2.add_paragraph()
    p.text = "• Severe global shortage of certified thoracic radiologists, causing diagnostic delays in rural/emergent settings.\n\n" \
             "• Fatigue and high workload induce inter-observer variability and missed subtle consolidations.\n\n" \
             "• Substantial visual overlap exists between viral/bacterial pneumonia and early COVID-19 manifestations."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(12)

    c3 = add_card(s2, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.0))
    tf_c3 = c3.text_frame
    tf_c3.word_wrap = True
    tf_c3.margin_left = tf_c3.margin_right = Inches(0.25)
    tf_c3.margin_top = Inches(0.3)
    p = tf_c3.paragraphs[0]
    p.text = "The 'Black-Box' Barrier"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = NAVY_DARK
    p = tf_c3.add_paragraph()
    p.text = "• Standard Deep Neural Networks output probability vectors without anatomical justification.\n\n" \
             "• Clinicians cannot ethically or legally trust opaque AI predictions without verifiable spatial evidence.\n\n" \
             "• Solution: Pair state-of-the-art ensemble accuracy with spatial Grad-CAM saliency maps."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(12)

    # ==================== SLIDE 3: Objectives & Scope ====================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Project Objectives & Scope Boundaries", "Target Deliverables & Scientific Rigor")

    c_obj = add_card(s3, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tf_obj = c_obj.text_frame
    tf_obj.word_wrap = True
    tf_obj.margin_left = tf_obj.margin_right = Inches(0.3)
    tf_obj.margin_top = Inches(0.3)
    p = tf_obj.paragraphs[0]
    p.text = "Key Project Objectives"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = NAVY_DARK
    p = tf_obj.add_paragraph()
    p.text = "1. Multi-Class Detection:\n" \
             "   Accurately categorize CXRs into 5 classes: Normal, COVID-19, Lung Opacity, Pneumonia, Tuberculosis.\n\n" \
             "2. Dual ResNet50 Ensembling:\n" \
             "   Combine independently trained models to minimize variance and improve boundary stability.\n\n" \
             "3. Interpretable Saliency Mapping:\n" \
             "   Generate pixel-aligned Grad-CAM heatmaps highlighting relevant anatomical regions.\n\n" \
             "4. Real-Time Clinical Web Platform:\n" \
             "   Deliver a sub-2-second responsive interface for triage, review, and diagnostic reporting."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(10)

    c_scope = add_card(s3, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tf_scope = c_scope.text_frame
    tf_scope.word_wrap = True
    tf_scope.margin_left = tf_scope.margin_right = Inches(0.3)
    tf_scope.margin_top = Inches(0.3)
    p = tf_scope.paragraphs[0]
    p.text = "Scope & Ethical Boundaries"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = NAVY_DARK
    p = tf_scope.add_paragraph()
    p.text = "• Academic Engineering Prototype:\n" \
             "   Designed as an educational and proof-of-concept research system, NOT a certified medical device.\n\n" \
             "• Assistive Decision-Support Tool:\n" \
             "   Supplements clinician judgment; does not replace board-certified radiologist evaluation.\n\n" \
             "• Standard Frontal CXR Modality:\n" \
             "   Focuses exclusively on standard adult frontal (AP/PA) views; lateral and CT modalities excluded.\n\n" \
             "• Ethical Disclosures:\n" \
             "   Complete audit of cross-split duplicate limits and patient identification constraints."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(10)

    # ==================== SLIDE 4: Proposed System Architecture ====================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Proposed System Architecture & End-to-End Pipeline", "Architectural Blueprint")

    boxes = [
        ("1. Input & Validation", "• Frontal CXR upload (PNG/JPG)\n• MIME type & integrity checks\n• OOD Saturation Gating\n• 224x224 RGB conversion", Inches(0.8)),
        ("2. Dual Model Ensemble", "• ResNet50 V1 + V3\n• Test-Time Augmentation (TTA)\n• Calibrated Clinical Cutoffs\n• P = 0.5 * (P_v1 + P_v3)", Inches(3.75)),
        ("3. Explainability Engine", "• Gradient Class Activation (Grad-CAM)\n• conv5_block3_3_conv gradients\n• Global Average Pooling + ReLU\n• Jet colormap overlay generation", Inches(6.7)),
        ("4. Clinical Frontend", "• React 18 + Vite Web Dashboard\n• Probability breakdown & alerts\n• Interactive heatmap overlay slider\n• History tracking & export", Inches(9.65))
    ]

    for title, desc, left in boxes:
        card = add_card(s4, left, Inches(1.8), Inches(2.85), Inches(4.8))
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.25)
        p = tf.paragraphs[0]
        p.text = title
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = SLATE_BLUE
        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(12)

    # ==================== SLIDE 5: Dataset Characterization ====================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Dataset Characterization & Disease Classes", "18,200 High-Resolution Radiographs")

    c_data_l = add_card(s5, Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.0))
    tf_dl = c_data_l.text_frame
    tf_dl.word_wrap = True
    tf_dl.margin_left = tf_dl.margin_right = Inches(0.3)
    tf_dl.margin_top = Inches(0.25)
    p = tf_dl.paragraphs[0]
    p.text = "Dataset Composition (N = 18,200 Images)"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = NAVY_DARK
    p = tf_dl.add_paragraph()
    p.text = "The study utilizes a curated collection of 18,200 frontal chest radiographs synthesized from verified open medical repositories (COVID-19 Radiography Database & TB Chest X-ray Database):\n\n" \
             "• Normal:  10,192 images (56.00%)\n" \
             "  Clear lung fields, sharp costophrenic angles, intact cardiac silhouette.\n\n" \
             "• COVID-19:  3,616 images (19.87%)\n" \
             "  Bilateral peripheral ground-glass opacities, consolidations.\n\n" \
             "• Lung Opacity:  2,347 images (12.90%)\n" \
             "  Non-specific pulmonary consolidations, fluid, or atelectasis.\n\n" \
             "• Pneumonia:  1,345 images (7.39%)\n" \
             "  Lobar consolidations, bacterial/viral bronchopneumonia.\n\n" \
             "• Tuberculosis:  700 images (3.85%)\n" \
             "  Apical infiltrates, cavitary lesions, fibrocalcific scars."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(8)

    c_data_r = add_card(s5, Inches(7.1), Inches(1.6), Inches(5.4), Inches(5.0))
    tf_dr = c_data_r.text_frame
    tf_dr.word_wrap = True
    tf_dr.margin_left = tf_dr.margin_right = Inches(0.3)
    tf_dr.margin_top = Inches(0.25)
    p = tf_dr.paragraphs[0]
    p.text = "Real-World Distribution Realism"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = NAVY_DARK
    p = tf_dr.add_paragraph()
    p.text = "• Clinical Imbalance Representation:\n" \
             "  Real-world hospitals evaluate mostly healthy patients, with fewer TB cases. Preserving this class hierarchy teaches the model realistic prior prevalence rather than artificial 20% uniform distributions.\n\n" \
             "• Rigorous Image Normalization:\n" \
             "  All heterogeneous source resolutions are standardized to 224x224 RGB.\n\n" \
             "• Class Separation:\n" \
             "  Distinguishes specific bacterial/viral pneumonia from non-specific opacities and COVID-19 manifestations."
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(10)

    # ==================== SLIDE 6: Preprocessing & Data Splits ====================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Data Preprocessing & Split Integrity", "Strict Manifest Partitioning")

    c_split = add_card(s6, Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.0))
    tf_sp = c_split.text_frame
    tf_sp.word_wrap = True
    tf_sp.margin_left = tf_sp.margin_right = Inches(0.3)
    tf_sp.margin_top = Inches(0.25)
    p = tf_sp.paragraphs[0]
    p.text = "Authoritative 70 / 15 / 15 Partitioning"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = NAVY_DARK
    p = tf_sp.add_paragraph()
    p.text = "The 18,200 active images are locked into persistent manifest JSON files to ensure exact scientific reproducibility:\n\n" \
             "• Training Manifest (70%):  12,740 images\n" \
             "  Used for feature representation and weight updates.\n\n" \
             "• Validation Manifest (15%):  2,730 images\n" \
             "  Used for hyperparameter tuning, checkpointing, and early stopping.\n\n" \
             "• Held-Out Test Manifest (15%):  2,730 images\n" \
             "  Completely unobserved during model training; dedicated to final statistical benchmark evaluation."
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(10)

    c_prep = add_card(s6, Inches(7.1), Inches(1.6), Inches(5.4), Inches(5.0))
    tf_pr = c_prep.text_frame
    tf_pr.word_wrap = True
    tf_pr.margin_left = tf_pr.margin_right = Inches(0.3)
    tf_pr.margin_top = Inches(0.25)
    p = tf_pr.paragraphs[0]
    p.text = "Pre-processing & Manifest Integrity"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = NAVY_DARK
    p = tf_pr.add_paragraph()
    p.text = "• Deterministic Resolution Rescaling:\n" \
             "  Bi-linear interpolation resizes all input matrices to 224x224x3.\n\n" \
             "• Channel Expansion:\n" \
             "  Monochrome radiographs are expanded to 3 identical RGB channels to match ImageNet pre-trained tensor dimensions.\n\n" \
             "• Floating-Point Scaling:\n" \
             "  Pixel intensities normalized to [0.0, 1.0] range.\n\n" \
             "• Manifest Immutability:\n" \
             "  Filepaths and target labels are locked in `dataset_manifests/` preventing data contamination."
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(10)

    # ==================== SLIDE 7: Dual-Model ResNet50 Ensemble ====================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Dual-Model ResNet50 Ensemble Architecture", "Variance Reduction via Deep Residual Learning")

    c_ens_l = add_card(s7, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tf_el = c_ens_l.text_frame
    tf_el.word_wrap = True
    tf_el.margin_left = tf_el.margin_right = Inches(0.3)
    tf_el.margin_top = Inches(0.25)
    p = tf_el.paragraphs[0]
    p.text = "Constituent Base Models"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = NAVY_DARK
    p = tf_el.add_paragraph()
    p.text = "• Why ResNet50?\n" \
             "  50-layer deep convolutional network with residual skip connections (F(x) + x) overcoming vanishing gradients and enabling rich thoracic feature hierarchies.\n\n" \
             "• Model V1 (final_super_model_old.keras):\n" \
             "  - Trained baseline ResNet50 architecture.\n" \
             "  - Standalone Test Accuracy: 96.08%.\n\n" \
             "• Model V3 (final_super_model_v3.keras):\n" \
             "  - Fine-tuned with modified bottleneck layers.\n" \
             "  - Standalone Test Accuracy: 97.18%."
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(10)

    c_ens_r = add_card(s7, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tf_er = c_ens_r.text_frame
    tf_er.word_wrap = True
    tf_er.margin_left = tf_er.margin_right = Inches(0.3)
    tf_er.margin_top = Inches(0.25)
    p = tf_er.paragraphs[0]
    p.text = "Production Ensemble Fusion"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = NAVY_DARK
    p = tf_er.add_paragraph()
    p.text = "• Softmax Probability Averaging:\n" \
             "  For each test sample x, predicted probability vector is computed as:\n" \
             "     P_ensemble(c|x) = 0.5 * P_v1(c|x) + 0.5 * P_v3(c|x)\n\n" \
             "• Test-Time Augmentation (TTA):\n" \
             "  Inference averages canonical and horizontal reflections for orientation invariance.\n\n" \
             "• Calibrated Decision Cutoffs:\n" \
             "  Tuned per class (TB: 0.35, Pneumonia: 0.38) preventing false-negative underdiagnosis.\n\n" \
             "• Ensemble Test Benchmark: 97.40% Accuracy\n" \
             "  Outperforms both individual models (96.08% and 97.18%), yielding 2,659 correct predictions out of 2,730."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(8)

    # ==================== SLIDE 8: Prediction Workflow & Tech Stack ====================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Prediction Workflow & Modern Technology Stack", "Full-Stack System Engineering")

    c_be = add_card(s8, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tf_be = c_be.text_frame
    tf_be.word_wrap = True
    tf_be.margin_left = tf_be.margin_right = Inches(0.3)
    tf_be.margin_top = Inches(0.25)
    p = tf_be.paragraphs[0]
    p.text = "High-Performance Backend"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = SLATE_BLUE
    p = tf_be.add_paragraph()
    p.text = "• Framework: FastAPI (Python 3.11) with ASGI Uvicorn.\n" \
             "• Deep Learning Engine: TensorFlow & Keras 2.21.\n" \
             "• Image Processing: OpenCV, Pillow, and NumPy.\n" \
             "• Asynchronous Concurrency:\n" \
             "  Model inferences and Grad-CAM generation dispatched to dedicated threadpools (`asyncio.to_thread`) ensuring the event loop remains responsive.\n" \
             "• Data Storage: MongoDB / Mongomock repository storing prediction logs, confidence distributions, and clinician review flags.\n" \
             "• Defensive Architecture: Strict Pydantic schema validation, JWT auth, and structured error boundaries."
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(10)

    c_fe = add_card(s8, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tf_fe = c_fe.text_frame
    tf_fe.word_wrap = True
    tf_fe.margin_left = tf_fe.margin_right = Inches(0.3)
    tf_fe.margin_top = Inches(0.25)
    p = tf_fe.paragraphs[0]
    p.text = "Responsive Frontend Dashboard"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = SLATE_BLUE
    p = tf_fe.add_paragraph()
    p.text = "• Framework: React 18 with Vite high-speed bundler.\n" \
             "• Visual Styling: Pure Vanilla CSS Design System with dark mode glassmorphism.\n" \
             "• Real-Time Interactive Canvas:\n" \
             "  Provides an opacity slider blending the original radiograph with the Grad-CAM heatmap dynamically.\n" \
             "• Clinical Review Workflow:\n" \
             "  Allows physicians to mark diagnoses as Confirmed, Flagged, or In-Review.\n" \
             "• Diagnostic Summary Export:\n" \
             "  One-click generation of comprehensive clinical triage reports."
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(10)

    # ==================== SLIDE 9: Grad-CAM Explainability ====================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Grad-CAM Explainability & Visual Saliency", "Localizing Pathological Features via Gradient Tracking")

    c_gc_l = add_card(s9, Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.0))
    tf_gcl = c_gc_l.text_frame
    tf_gcl.word_wrap = True
    tf_gcl.margin_left = tf_gcl.margin_right = Inches(0.3)
    tf_gcl.margin_top = Inches(0.2)
    p = tf_gcl.paragraphs[0]
    p.text = "Mathematical Formulation & Pipeline"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = NAVY_DARK
    p = tf_gcl.add_paragraph()
    p.text = "1. Target Activation Layer:\n" \
             "   Layer `conv5_block3_3_conv` — the final convolutional feature block in ResNet50 (spatial resolution 7x7x2048).\n\n" \
             "2. Gradient Computation:\n" \
             "   Calculates gradient of predicted score y^c w.r.t feature map activations A^k:  d(y^c) / d(A^k)\n\n" \
             "3. Importance Weights (alpha_k^c):\n" \
             "   Global Average Pooling computes importance of each filter.\n\n" \
             "4. Rectified Combination (ReLU):\n" \
             "   L_Grad-CAM = ReLU( sum_k alpha_k^c * A^k )\n" \
             "   Isolates features that positively contribute to class c.\n\n" \
             "• Critical Limitation Disclosed:\n" \
             "  Grad-CAM visualizes Model V3 specifically, not the full ensemble."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(6)

    # Add real Grad-CAM screenshot image on right
    img_gc_path = r"D:\7TH SEM MAJOR F\Frontend chest x ray\assets\screenshots\05_gradcam_explainability.png"
    if os.path.exists(img_gc_path):
        s9.shapes.add_picture(img_gc_path, Inches(7.1), Inches(1.6), width=Inches(5.4))
        # Add caption below picture
        cap = s9.shapes.add_textbox(Inches(7.1), Inches(5.8), Inches(5.4), Inches(0.8))
        tf_cap = cap.text_frame
        tf_cap.word_wrap = True
        p_c = tf_cap.paragraphs[0]
        p_c.text = "Verified Application Screenshot: Bilateral infiltration localization with interactive opacity control and probability breakdown."
        p_c.font.size = Pt(10)
        p_c.font.color.rgb = TEXT_MUTED

    # ==================== SLIDE 10: Evaluation Methodology & Metrics ====================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Evaluation Methodology & Comprehensive Metrics", "Statistical Validation on N=2,730 Test Radiographs")

    # 4 metric cards across top
    metrics_summary = [
        ("97.40%", "Ensemble Accuracy", "2,659 / 2,730 Correct", Inches(0.8)),
        ("0.9668", "Macro F1-Score", "Unweighted class mean", Inches(3.75)),
        ("0.9740", "Weighted F1-Score", "Support-weighted mean", Inches(6.7)),
        ("0.9986", "Macro ROC-AUC", "Near-perfect discrimination", Inches(9.65))
    ]
    for val, lbl, sub, left in metrics_summary:
        c_m = add_card(s10, left, Inches(1.6), Inches(2.85), Inches(1.6))
        tf = c_m.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.15)
        tf.margin_top = Inches(0.15)
        p = tf.paragraphs[0]
        p.text = val
        p.font.bold = True
        p.font.size = Pt(24)
        p.font.color.rgb = TEAL_ACCENT
        p.font.name = "Arial"
        p = tf.add_paragraph()
        p.text = lbl
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = NAVY_DARK
        p = tf.add_paragraph()
        p.text = sub
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_MUTED

    # Detailed evaluation methodology card below
    c_eval = add_card(s10, Inches(0.8), Inches(3.5), Inches(11.733), Inches(3.2))
    tf_ev = c_eval.text_frame
    tf_ev.word_wrap = True
    tf_ev.margin_left = tf_ev.margin_right = Inches(0.3)
    tf_ev.margin_top = Inches(0.25)
    p = tf_ev.paragraphs[0]
    p.text = "Authoritative Evaluation Standards (Derived from metrics.json)"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = NAVY_DARK
    p = tf_ev.add_paragraph()
    p.text = "• Macro-Averaging vs Weighted-Averaging:\n" \
             "  Macro metrics evaluate all 5 diseases equally regardless of sample frequency, proving the system does not inflate accuracy by merely favoring the majority 'Normal' class. The Macro F1 of 0.9668 confirms high diagnostic balance.\n\n" \
             "• Standalone Model Comparison:\n" \
             "  - Model V1 (Original ResNet50):  Accuracy = 96.08%  (2,623 / 2,730 correct)\n" \
             "  - Model V3 (Fine-tuned ResNet50): Accuracy = 97.18%  (2,653 / 2,730 correct)\n" \
             "  - Ensemble (V1 + V3 Average):    Accuracy = 97.40%  (2,659 / 2,730 correct) — best across all metrics."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(8)

    # ==================== SLIDE 11: Class-Wise Results & Curves ====================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "Verified Class-Wise Performance & ROC Analysis", "Zero False-Negative Lung Opacities & Robust Curves")

    # Table card on left
    c_tbl = add_card(s11, Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.0))
    tf_tb = c_tbl.text_frame
    tf_tb.word_wrap = True
    tf_tb.margin_left = tf_tb.margin_right = Inches(0.25)
    tf_tb.margin_top = Inches(0.2)
    p = tf_tb.paragraphs[0]
    p.text = "Class-Wise Metrics Table (Test N=2,730)"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = NAVY_DARK
    p = tf_tb.add_paragraph()
    p.text = "• Lung Opacity (Support: 352):\n" \
             "  Precision: 100.00% | Recall: 100.00% | F1: 100.00%\n\n" \
             "• Normal (Support: 1,529):\n" \
             "  Precision: 98.23% | Recall: 97.97% | F1: 98.10%\n\n" \
             "• Pneumonia (Support: 201):\n" \
             "  Precision: 96.53% | Recall: 97.01% | F1: 96.77%\n\n" \
             "• COVID-19 (Support: 543):\n" \
             "  Precision: 94.50% | Recall: 94.84% | F1: 94.67%\n\n" \
             "• Tuberculosis (Support: 105):\n" \
             "  Precision: 93.40% | Recall: 94.29% | F1: 93.84%"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(8)

    # Confusion matrix & ROC Curves pictures on right
    cm_path = r"D:\7TH SEM MAJOR F\Backend chest x ray\evaluation\confusion_matrix.png"
    roc_path = r"D:\7TH SEM MAJOR F\Backend chest x ray\evaluation\roc_auc_curves.png"
    if os.path.exists(cm_path):
        s11.shapes.add_picture(cm_path, Inches(6.9), Inches(1.6), width=Inches(2.8))
    if os.path.exists(roc_path):
        s11.shapes.add_picture(roc_path, Inches(9.85), Inches(1.6), width=Inches(2.8))

    cap_curves = s11.shapes.add_textbox(Inches(6.9), Inches(5.8), Inches(5.7), Inches(0.8))
    tf_cc = cap_curves.text_frame
    tf_cc.word_wrap = True
    p_cc = tf_cc.paragraphs[0]
    p_cc.text = "Figure 11.1: Official Confusion Matrix (left) and Multi-Class ROC Curves (right), displaying near-perfect area under curve (Macro ROC-AUC = 0.9986)."
    p_cc.font.size = Pt(10)
    p_cc.font.color.rgb = TEXT_MUTED

    # ==================== SLIDE 12: Runtime Performance & CPU Constraints ====================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "Runtime Performance & Computational Profiling", "Benchmarked on Local CPU Hardware")

    c_perf1 = add_card(s12, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tf_p1 = c_perf1.text_frame
    tf_p1.word_wrap = True
    tf_p1.margin_left = tf_p1.margin_right = Inches(0.3)
    tf_p1.margin_top = Inches(0.25)
    p = tf_p1.paragraphs[0]
    p.text = "Local Hardware Benchmark Profile"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = NAVY_DARK
    p = tf_p1.add_paragraph()
    p.text = "Audited on local Intel Core CPU execution environment (without GPU acceleration):\n\n" \
             "• Image Decoding & Preprocessing:  ~30 ms\n" \
             "  Resizing, byte parsing, channel replication.\n\n" \
             "• Dual ResNet50 Ensemble Inference:  ~675 ms\n" \
             "  Sequential forward passes of V1 + V3 models and probability vector averaging.\n\n" \
             "• Grad-CAM Saliency Computation:  ~825 ms\n" \
             "  Backpropagating gradients to conv5_block3_3_conv and alpha blending heatmap.\n\n" \
             "• Total Processing Turnaround:  ~1.50 seconds\n" \
             "  Exceeds standard 2-second clinical responsiveness requirement."
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(10)

    c_perf2 = add_card(s12, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tf_p2 = c_perf2.text_frame
    tf_p2.word_wrap = True
    tf_p2.margin_left = tf_p2.margin_right = Inches(0.3)
    tf_p2.margin_top = Inches(0.25)
    p = tf_p2.paragraphs[0]
    p.text = "CPU Optimization Strategies"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = NAVY_DARK
    p = tf_p2.add_paragraph()
    p.text = "• Model Pre-Warming:\n" \
             "  Models are loaded and compiled during FastAPI startup. Zero disk I/O penalty occurs during incoming HTTP requests.\n\n" \
             "• Asynchronous Thread Offloading:\n" \
             "  CPU-heavy matrix operations execute inside Python's worker threads (`asyncio.to_thread`), preventing server lockup.\n\n" \
             "• In-Memory Heatmap Streaming:\n" \
             "  Grad-CAM heatmaps are converted directly to base64 data URIs in RAM without writing temporary files to storage.\n\n" \
             "• Future GPU Acceleration:\n" \
             "  With NVIDIA CUDA deployment, inference latency will drop below 60 ms."
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(10)

    # ==================== SLIDE 13: Application UI Showcase ====================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_background(s13)
    add_header(s13, "Application Showcase & Clinical User Journey", "Genuine Screenshots of Verified Full-Stack Prototype")

    # Embed 2 real screenshots side by side: dashboard & results
    img_dash_path = r"D:\7TH SEM MAJOR F\Frontend chest x ray\assets\screenshots\02_main_dashboard.png"
    img_res_path = r"D:\7TH SEM MAJOR F\Frontend chest x ray\assets\screenshots\04_prediction_results.png"
    
    if os.path.exists(img_dash_path):
        s13.shapes.add_picture(img_dash_path, Inches(0.8), Inches(1.6), width=Inches(5.7))
        cap1 = s13.shapes.add_textbox(Inches(0.8), Inches(5.6), Inches(5.7), Inches(0.8))
        tf1 = cap1.text_frame
        tf1.word_wrap = True
        p1 = tf1.paragraphs[0]
        p1.text = "Dashboard: System model status, class metrics summary, live latency counters, and quick triage upload."
        p1.font.size = Pt(10.5)
        p1.font.color.rgb = TEXT_DARK

    if os.path.exists(img_res_path):
        s13.shapes.add_picture(img_res_path, Inches(6.8), Inches(1.6), width=Inches(5.7))
        cap2 = s13.shapes.add_textbox(Inches(6.8), Inches(5.6), Inches(5.7), Inches(0.8))
        tf2 = cap2.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = "Diagnostic View: Predicted class confidence, ranked disease probabilities, and doctor review triage actions."
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_DARK

    # ==================== SLIDE 14: Research Limitations ====================
    s14 = prs.slides.add_slide(blank_layout)
    set_slide_background(s14)
    add_header(s14, "Research Limitations & Scientific Transparency", "Disclosed Vulnerabilities & Rigorous Audit")

    alerts = [
        ("Nine Exact Cross-Split Duplicate Images (9 Pairs)", 
         "A byte-for-byte MD5 hash audit identified nine exact duplicate images (9 pairs) shared between the training (12,740) and testing (2,730) manifests (0.049%). While minimal, transparent scientific reporting acknowledges this potential source of evaluation bias.",
         Inches(1.6)),
        ("Absence of Patient Identifiers", 
         "Publicly available CXR databases omit patient identifiers. As a result, split independence is image-level rather than patient-level, meaning multiple radiographs from the same patient could potentially exist across splits.",
         Inches(3.3)),
        ("Grad-CAM Architectural Disconnect & Non-Clinical Status", 
         "Grad-CAM computes saliency with respect to Model V3 specifically rather than the joint mathematical ensemble. Additionally, this system is strictly an academic prototype and is NOT clinically validated for medical diagnosis.",
         Inches(5.0))
    ]

    for title, desc, top in alerts:
        c_al = add_card(s14, Inches(0.8), top, Inches(11.733), Inches(1.5), bg_color=ALERT_BG, border_color=ALERT_BORDER)
        tf = c_al.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.18)
        p = tf.paragraphs[0]
        p.text = f"• {title}"
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = ALERT_TEXT
        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(4)

    # ==================== SLIDE 15: Conclusion & Future Scope ====================
    s15 = prs.slides.add_slide(blank_layout)
    set_slide_background(s15)
    add_header(s15, "Conclusions & Future Research Directions", "Summary of Contributions & Roadmap Ahead")

    c_concl = add_card(s15, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tf_co = c_concl.text_frame
    tf_co.word_wrap = True
    tf_co.margin_left = tf_co.margin_right = Inches(0.3)
    tf_co.margin_top = Inches(0.25)
    p = tf_co.paragraphs[0]
    p.text = "Key Project Achievements"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = NAVY_DARK
    p = tf_co.add_paragraph()
    p.text = "• 97.40% Verified Ensemble Accuracy:\n" \
             "  Demonstrated strong classification across 2,730 unseen test images with 0.9668 Macro F1 and 0.9986 ROC-AUC.\n\n" \
             "• Transparent Explainability:\n" \
             "  Coupled predictions with visual Grad-CAM heatmaps to localize pathological consolidations.\n\n" \
             "• Complete Full-Stack Solution:\n" \
             "  Built a responsive, production-built web application with doctor review actions and history tracking.\n\n" \
             "• Academic Honesty:\n" \
             "  Full disclosure of duplicates, patient ID limitations, and architectural trade-offs."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(8)

    c_fut = add_card(s15, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tf_fu = c_fut.text_frame
    tf_fu.word_wrap = True
    tf_fu.margin_left = tf_fu.margin_right = Inches(0.3)
    tf_fu.margin_top = Inches(0.25)
    p = tf_fu.paragraphs[0]
    p.text = "Future Technical Roadmap"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = NAVY_DARK
    p = tf_fu.add_paragraph()
    p.text = "• Hospital PACS & DICOM Integration:\n" \
             "  Direct ingestion of 16-bit DICOM digital radiographs from hospital picture archiving systems.\n\n" \
             "• True Ensemble Saliency Explanations:\n" \
             "  Implement Integrated Gradients or Attention Rollout across both models simultaneously.\n\n" \
             "• Multi-Center Clinical Trials:\n" \
             "  Prospective institutional validation on hospital-acquired cohorts with strict patient identifiers.\n\n" \
             "• Multi-Modal AI:\n" \
             "  Integrate patient vitals, lab reports, and demographic clinical data with image representations."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(8)

    output_path = r"D:\7TH SEM MAJOR F\presentation.pptx"
    prs.save(output_path)
    print(f"Presentation generated successfully: {output_path}")
    print(f"Total slides: {len(prs.slides)}")

if __name__ == "__main__":
    create_deck()
